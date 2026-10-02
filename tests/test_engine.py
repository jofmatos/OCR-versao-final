"""Behavior checks use real PDFs and, when installed, the real OCR executable."""

import io
import shutil
import threading
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pymupdf as fitz
import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont

from app import engine


@pytest.fixture
def native_pdf(tmp_path):
    path = tmp_path / "texto.pdf"
    with fitz.open() as document:
        first = document.new_page()
        first.insert_text((60, 80), "Conversão de documentos com acentuação.", fontsize=14)
        first.insert_text((60, 130), "Segunda linha: informação e revisão.", fontsize=12)
        second = document.new_page()
        second.insert_text((60, 80), "Página dois. Conteúdo editável.", fontsize=14)
        document.save(path)
    return path


def test_native_pdf_keeps_accents_order_and_progress(native_pdf):
    callbacks = []
    result = engine.convert_document(native_pdf, page_numbers=[2, 1], on_page=callbacks.append)
    assert [page["number"] for page in result] == [2, 1]
    assert callbacks == result
    assert all(page["method"] == "native" for page in result)
    assert "Conteúdo editável" in result[0]["text"]
    assert "acentuação" in result[1]["text"]
    assert "revisão" in result[1]["text"]
    assert engine.inspect_document(native_pdf)["page_count"] == 2


def test_preview_is_png_and_bounded(native_pdf):
    with Image.open(io.BytesIO(engine.render_preview(native_pdf, 1))) as preview:
        assert preview.format == "PNG"
        assert max(preview.size) <= 1400


def test_invalid_selected_page_fails_before_callbacks(native_pdf):
    callbacks = []
    with pytest.raises(engine.ProcessingError, match="intervalo"):
        engine.convert_document(native_pdf, page_numbers=[1, 3], on_page=callbacks.append)
    assert callbacks == []


def test_encrypted_pdf_has_actionable_error(tmp_path):
    path = tmp_path / "senha.pdf"
    with fitz.open() as document:
        document.new_page()
        document.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="secret")
    with pytest.raises(engine.ProcessingError, match="senha"):
        engine.inspect_document(path)


def test_corrupt_file_has_actionable_error(tmp_path):
    path = tmp_path / "quebrado.pdf"
    path.write_bytes(b"this is not a PDF")
    with pytest.raises(engine.ProcessingError, match="PDF"):
        engine.inspect_document(path)


def test_cancellation_between_pages_keeps_first_callback(native_pdf):
    callbacks = []
    with pytest.raises(engine.ProcessingCancelled):
        engine.convert_document(native_pdf, page_numbers=[1, 2], on_page=callbacks.append, cancelled=lambda: bool(callbacks))
    assert [page["number"] for page in callbacks] == [1]


def test_blank_page_is_not_sent_to_ocr(tmp_path, monkeypatch):
    path = tmp_path / "vazio.pdf"
    with fitz.open() as document:
        document.new_page()
        document.save(path)
    monkeypatch.setattr(engine, "_validate_languages", lambda language: pytest.fail("Blank native page must not require Tesseract"))
    result = engine.convert_document(path, page_numbers=[1])
    assert result[0]["text"] == ""
    assert result[0]["method"] == "native"


def test_missing_languages_are_not_silently_replaced(monkeypatch):
    monkeypatch.setattr(engine, "get_ocr_languages", lambda: ["eng"])
    with pytest.raises(engine.ProcessingError, match="por"):
        engine._validate_languages("por+eng")
    with pytest.raises(engine.ProcessingError, match="inválida"):
        engine._validate_languages("eng --psm 0")


def test_docx_and_txt_contain_editable_accented_text_and_page_breaks(native_pdf):
    pages = engine.convert_document(native_pdf, page_numbers=[1, 2])
    assert "acentuação" in engine.export_txt(pages).decode("utf-8")
    document = Document(io.BytesIO(engine.export_docx(pages, "Relatório")))
    assert document.core_properties.title == "Relatório"
    assert "acentuação" in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Conteúdo editável" in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert len(document.element.xpath('.//w:br[@w:type="page"]')) == 1


def test_docx_filters_unsupported_xml_characters():
    output = engine.export_docx([{"text": "texto\x00 com acento: ação\ud800"}], "título\x00")
    document = Document(io.BytesIO(output))
    assert document.paragraphs[0].text == "texto com acento: ação"


def _scanned_pdf(path: Path, *, header: bool = False):
    font_path = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    if not font_path.exists():
        pytest.skip("The integration fixture needs DejaVu Sans")
    image = Image.new("RGB", (1654, 2339), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(font_path), 42)
    lines = [
        "Conversão de documentos em português.",
        "Informação, educação e comunicação.",
        "Este documento permite revisão do texto.",
        "O relatório contém números: 12345.",
        "Uma transcrição deve preservar acentuação.",
    ]
    for index, line in enumerate(lines):
        draw.text((120, 300 + index * 85), line, font=font, fill="black")
    output = io.BytesIO()
    image.save(output, "PNG")
    with fitz.open() as document:
        page = document.new_page(width=595, height=842)
        page.insert_image(page.rect, stream=output.getvalue())
        if header:
            page.insert_text((40, 35), "Arquivo 1", fontsize=10)
        document.save(path)


def test_auto_ocr_reads_real_scan_even_with_native_header(tmp_path):
    if not shutil.which("tesseract") or "por" not in engine.get_ocr_languages():
        pytest.skip("Integration OCR requires Tesseract and Portuguese language data")
    path = tmp_path / "digitalizado.pdf"
    _scanned_pdf(path, header=True)
    result = engine.convert_document(path, page_numbers=[1], language="por+eng", quality="standard")
    assert result[0]["method"] == "ocr"
    assert result[0]["confidence"] > 70
    for phrase in ["português", "Informação", "12345", "acentuação"]:
        assert phrase in result[0]["text"]


def test_forced_ocr_of_native_pdf(native_pdf):
    if not shutil.which("tesseract") or "por" not in engine.get_ocr_languages():
        pytest.skip("Integration OCR requires Tesseract and Portuguese language data")
    result = engine.convert_document(native_pdf, page_numbers=[1], mode="ocr", quality="standard")
    assert result[0]["method"] == "ocr"
    assert "documentos" in result[0]["text"]


def test_preview_remains_available_while_ocr_is_running(native_pdf, monkeypatch):
    started = threading.Event()
    finish = threading.Event()

    def delayed_ocr(*args):
        started.set()
        assert finish.wait(timeout=5)
        return "texto", 99, []

    monkeypatch.setattr(engine, "_transcribe", delayed_ocr)
    monkeypatch.setattr(engine, "_validate_languages", lambda language: None)
    with ThreadPoolExecutor(max_workers=2) as executor:
        conversion = executor.submit(engine.convert_document, native_pdf, page_numbers=[1], mode="ocr")
        try:
            assert started.wait(timeout=3)
            preview = executor.submit(engine.render_preview, native_pdf, 2)
            assert preview.result(timeout=2).startswith(b"\x89PNG")
        finally:
            finish.set()
        assert conversion.result(timeout=3)[0]["text"] == "texto"


def test_real_ocr_corrects_rotated_scans(tmp_path):
    if not shutil.which("tesseract") or "por" not in engine.get_ocr_languages():
        pytest.skip("Integration OCR requires Tesseract and Portuguese language data")
    original = tmp_path / "original.pdf"
    rotated = tmp_path / "girado.pdf"
    _scanned_pdf(original)
    with fitz.open(original) as document:
        document[0].set_rotation(90)
        document.save(rotated)
    result = engine.convert_document(rotated, page_numbers=[1], quality="high")
    assert "acentuação" in result[0]["text"]
    assert any("Orientação corrigida" in warning for warning in result[0]["warnings"])
