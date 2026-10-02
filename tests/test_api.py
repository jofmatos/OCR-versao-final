"""Exercise the document workflow through HTTP with real PDF/DOCX files."""

from io import BytesIO
import json
from threading import Event
import time

from docx import Document
from fastapi.testclient import TestClient
from PIL import Image
import pymupdf as fitz
import pytest

from app import engine
from app.main import create_app


@pytest.fixture
def pdf_bytes():
    with fitz.open() as document:
        first = document.new_page()
        first.insert_text((60, 80), "Informação e educação: revisão com acentuação.", fontsize=14)
        second = document.new_page()
        second.insert_text((60, 80), "Página dois: conteúdo editável em português.", fontsize=14)
        return document.tobytes()


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path / "documents")) as connection:
        yield connection


@pytest.fixture
def paused_conversion(monkeypatch):
    entered = Event()
    release = Event()
    real_convert = engine.convert_document

    def slow_conversion(*args, **kwargs):
        entered.set()
        if not release.wait(10):
            raise RuntimeError("Test did not release the worker")
        return real_convert(*args, **kwargs)

    monkeypatch.setattr(engine, "convert_document", slow_conversion)
    try:
        yield entered, release
    finally:
        release.set()


def upload(client, content, name="Relatório.pdf"):
    response = client.post("/api/documents", files={"file": (name, content, "application/pdf")})
    assert response.status_code == 201, response.text
    return response.json()


def wait_for_result(client, identifier):
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        response = client.get(f"/api/documents/{identifier}")
        assert response.status_code == 200, response.text
        result = response.json()
        if result["status"] != "processing":
            assert result["status"] == "ready", result
            return result
        time.sleep(0.02)
    pytest.fail("Document conversion did not finish within 10 seconds")


def convert(client, identifier, **options):
    response = client.post(f"/api/documents/{identifier}/convert", json=options)
    assert response.status_code == 202, response.text
    return wait_for_result(client, identifier)


def test_pdf_to_preview_editable_docx_and_utf8_text(client, pdf_bytes):
    record = upload(client, pdf_bytes)
    identifier = record["id"]
    assert record["name"] == "Relatório.pdf"
    assert record["page_count"] == 2
    assert record["size"] == len(pdf_bytes)
    assert record["status"] == "uploaded"

    preview = client.get(f"/api/documents/{identifier}/preview/2")
    assert preview.status_code == 200
    assert preview.headers["content-type"] == "image/png"
    with Image.open(BytesIO(preview.content)) as image:
        assert image.format == "PNG"
        assert 0 < max(image.size) <= 1400
    assert client.get(f"/api/documents/{identifier}/preview/3").status_code == 404
    assert client.get(f"/api/documents/{identifier}/export").status_code == 409

    ready = convert(client, identifier)
    assert ready["progress"] == {"completed": 2, "total": 2}
    assert [page["number"] for page in ready["pages"]] == [1, 2]
    assert all(page["method"] == "native" for page in ready["pages"])
    assert "acentuação" in ready["pages"][0]["text"]
    assert "português" in ready["pages"][1]["text"]

    txt = client.get(f"/api/documents/{identifier}/export?format=txt")
    assert txt.status_code == 200
    assert "utf-8" in txt.headers["content-type"]
    assert "Informação" in txt.content.decode("utf-8")
    assert txt.text.index("Informação") < txt.text.index("Página dois")
    assert "Relat%C3%B3rio.txt" in txt.headers["content-disposition"]

    docx = client.get(f"/api/documents/{identifier}/export?format=docx")
    assert docx.status_code == 200
    assert docx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    document = Document(BytesIO(docx.content))
    assert document.core_properties.title == "Relatório"
    paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Informação" in paragraphs
    assert "conteúdo editável" in paragraphs
    assert len(document.element.xpath('.//w:br[@w:type="page"]')) == 1


def test_edits_survive_restart_and_are_exported(tmp_path, pdf_bytes):
    directory = tmp_path / "persistent"
    corrected = "Correção revisada: ação, órgãos e informações.\n\nNovo parágrafo."
    with TestClient(create_app(directory)) as client:
        identifier = upload(client, pdf_bytes)["id"]
        original = convert(client, identifier)
        response = client.patch(f"/api/documents/{identifier}/text", json={"pages": [{"number": 1, "text": corrected}]})
        assert response.status_code == 200, response.text
        assert response.json()["pages"][1]["text"] == original["pages"][1]["text"]

    with TestClient(create_app(directory)) as client:
        restored = client.get(f"/api/documents/{identifier}").json()
        assert restored["status"] == "ready"
        assert restored["pages"][0]["text"] == corrected
        txt = client.get(f"/api/documents/{identifier}/export?format=txt")
        assert corrected in txt.content.decode("utf-8")
        docx = Document(BytesIO(client.get(f"/api/documents/{identifier}/export?format=docx").content))
        assert corrected.split("\n\n")[0] in [paragraph.text for paragraph in docx.paragraphs]
        assert "Novo parágrafo." in [paragraph.text for paragraph in docx.paragraphs]
        assert client.delete(f"/api/documents/{identifier}").status_code == 204
        assert client.get(f"/api/documents/{identifier}").status_code == 404
        assert client.get(f"/api/documents/{identifier}/preview/1").status_code == 404
        assert not (directory / identifier).exists()


@pytest.mark.parametrize("selection", ["0", "3", "2-1", "1-3", "1,,2", "1;2", "-1", "1-", "one", "1, 3"])
def test_invalid_page_selection_does_not_start_work(client, pdf_bytes, selection):
    identifier = upload(client, pdf_bytes)["id"]
    response = client.post(f"/api/documents/{identifier}/convert", json={"pages": selection})
    assert response.status_code == 400
    assert client.get(f"/api/documents/{identifier}").json()["status"] == "uploaded"


@pytest.mark.parametrize(("selection", "expected"), [("2", [2]), (" 2, 1-2, 1 ", [1, 2]), (" ", [1, 2])])
def test_selected_pages_are_exported_once_in_document_order(client, pdf_bytes, selection, expected):
    identifier = upload(client, pdf_bytes)["id"]
    result = convert(client, identifier, pages=selection)
    assert [page["number"] for page in result["pages"]] == expected
    assert result["progress"] == {"completed": len(expected), "total": len(expected)}
    output = client.get(f"/api/documents/{identifier}/export?format=txt").text
    assert ("Informação" in output) == (1 in expected)
    assert output.count("Página dois") == 1


@pytest.mark.parametrize(("content", "status"), [(b"", 400), (b"not a PDF", 415), (b"%PDF-1.7\ncorrupted data", 400)])
def test_invalid_uploads_leave_no_files(client, tmp_path, content, status):
    response = client.post("/api/documents", files={"file": ("invalid.pdf", content, "application/pdf")})
    assert response.status_code == status
    assert list((tmp_path / "documents").iterdir()) == []


def test_encrypted_upload_has_actionable_error_and_leaves_no_files(client, tmp_path):
    with fitz.open() as document:
        document.new_page()
        encrypted = document.tobytes(encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="secret")
    response = client.post("/api/documents", files={"file": ("encrypted.pdf", encrypted, "application/pdf")})
    assert response.status_code == 400
    assert "senha" in response.json()["detail"]
    assert list((tmp_path / "documents").iterdir()) == []


def test_processing_blocks_duplicate_jobs_edits_and_downloads(client, pdf_bytes, paused_conversion):
    entered, release = paused_conversion
    identifier = upload(client, pdf_bytes)["id"]
    try:
        assert client.post(f"/api/documents/{identifier}/convert", json={}).status_code == 202
        assert entered.wait(5)
        assert client.post(f"/api/documents/{identifier}/convert", json={}).status_code == 409
        assert client.patch(f"/api/documents/{identifier}/text", json={"pages": []}).status_code == 409
        assert client.get(f"/api/documents/{identifier}/export").status_code == 409
    finally:
        release.set()
    assert wait_for_result(client, identifier)["page_count"] == 2


def test_deleted_active_job_cannot_reappear_after_restart(client, tmp_path, pdf_bytes, paused_conversion):
    entered, release = paused_conversion
    directory = tmp_path / "documents"
    identifier = upload(client, pdf_bytes)["id"]
    stored = directory / identifier
    try:
        assert client.post(f"/api/documents/{identifier}/convert", json={}).status_code == 202
        assert entered.wait(5)
        assert (stored / "state.json").exists()
        assert client.delete(f"/api/documents/{identifier}").status_code == 204
        assert client.get(f"/api/documents/{identifier}").status_code == 404
        # An in-flight worker still owns the PDF, but deletion must already be
        # durable if the server stops before that worker finishes.
        assert not (stored / "state.json").exists()
        with TestClient(create_app(directory)) as restarted:
            assert restarted.get(f"/api/documents/{identifier}").status_code == 404
    finally:
        release.set()
    deadline = time.monotonic() + 5
    while stored.exists() and time.monotonic() < deadline:
        time.sleep(0.02)
    assert not stored.exists()


@pytest.mark.parametrize("pages", [[{"number": 3, "text": "invalid"}], [{"number": 1, "text": "first"}, {"number": 1, "text": "duplicate"}], [{"number": 1, "text": "unsafe\u0000text"}]])
def test_invalid_edits_preserve_previous_text(client, pdf_bytes, pages):
    identifier = upload(client, pdf_bytes)["id"]
    original = convert(client, identifier)
    response = client.patch(f"/api/documents/{identifier}/text", json={"pages": pages})
    assert response.status_code == 400
    assert client.get(f"/api/documents/{identifier}").json()["pages"] == original["pages"]


@pytest.mark.parametrize(("character", "status"), [("\x00", 400), ("\x0b", 400), ("\ud800", 422), ("\udfff", 422), ("\ufffe", 400), ("\uffff", 400)], ids=["null", "vertical-tab", "high-surrogate", "low-surrogate", "noncharacter-fffe", "noncharacter-ffff"])
def test_invalid_unicode_edit_is_rejected_without_damaging_exports(client, pdf_bytes, character, status):
    identifier = upload(client, pdf_bytes)["id"]
    original = convert(client, identifier)
    # JSON escapes let lone surrogates reach the server; directly encoding
    # these strings as UTF-8 would fail inside the client before HTTP begins.
    payload = json.dumps({"pages": [{"number": 1, "text": "Revisão " + character}]})
    response = client.patch(f"/api/documents/{identifier}/text", content=payload.encode("ascii"), headers={"Content-Type": "application/json"})
    assert response.status_code == status
    assert client.get(f"/api/documents/{identifier}").json()["pages"] == original["pages"]
    assert "Informação" in client.get(f"/api/documents/{identifier}/export?format=txt").content.decode("utf-8")
    exported = client.get(f"/api/documents/{identifier}/export?format=docx")
    assert exported.status_code == 200
    assert "Informação" in "\n".join(paragraph.text for paragraph in Document(BytesIO(exported.content)).paragraphs)


def test_interrupted_job_can_be_retried_after_restart(tmp_path, pdf_bytes):
    directory = tmp_path / "recovery"
    with TestClient(create_app(directory)) as client:
        identifier = upload(client, pdf_bytes)["id"]
    state_path = directory / identifier / "state.json"
    state = json.loads(state_path.read_text("utf-8"))
    state["status"] = "processing"
    state_path.write_text(json.dumps(state), encoding="utf-8")
    with TestClient(create_app(directory)) as client:
        recovered = client.get(f"/api/documents/{identifier}").json()
        assert recovered["status"] == "error"
        assert "reiniciado" in recovered["error"]
        assert convert(client, identifier)["status"] == "ready"


def test_expired_files_are_removed_on_startup(tmp_path, pdf_bytes):
    directory = tmp_path / "expiration"
    with TestClient(create_app(directory)) as client:
        identifier = upload(client, pdf_bytes)["id"]
    state_path = directory / identifier / "state.json"
    state = json.loads(state_path.read_text("utf-8"))
    state["updated_at"] = 0
    state_path.write_text(json.dumps(state), encoding="utf-8")
    with TestClient(create_app(directory)) as client:
        assert client.get(f"/api/documents/{identifier}").status_code == 404
        assert not (directory / identifier).exists()


def test_external_sites_cannot_upload_or_delete_documents(client, pdf_bytes):
    origin = {"Origin": "https://unrelated.example"}
    response = client.post("/api/documents", files={"file": ("pdf.pdf", pdf_bytes, "application/pdf")}, headers=origin)
    assert response.status_code == 403
    identifier = upload(client, pdf_bytes)["id"]
    assert client.delete(f"/api/documents/{identifier}", headers=origin).status_code == 403
    assert client.get(f"/api/documents/{identifier}").status_code == 200
    assert client.delete(f"/api/documents/{identifier}", headers={"Origin": "http://testserver"}).status_code == 204


def test_health_exposes_capabilities_and_api_disables_caching(client):
    response = client.get("/api/health")
    assert response.status_code == 200
    status = response.json()
    assert status["status"] == "ok"
    assert isinstance(status["ocr_available"], bool)
    assert isinstance(status["languages"], list)
    assert status["limits"] == {"max_upload_mb": 50, "max_pages": 200}
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"


def test_oversized_request_is_rejected_before_storage(client, pdf_bytes, tmp_path):
    response = client.post("/api/documents", files={"file": ("large.pdf", pdf_bytes, "application/pdf")}, headers={"Content-Length": str(52 * 1024 * 1024)})
    assert response.status_code == 413
    assert list((tmp_path / "documents").iterdir()) == []
